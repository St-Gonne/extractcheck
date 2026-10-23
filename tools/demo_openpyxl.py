"""Compare two openpyxl reading modes on one generated workbook only."""
from __future__ import annotations

import argparse
import hashlib
import json
import os
from pathlib import Path
import zipfile

from extractcheck.census import census_file
from extractcheck.contracts import validate_unit
from extractcheck.safety import atomic_json, deny_python_network, secure_directory, source_read_guard


def build_workbook(path: Path) -> None:
    """Write a deterministic OOXML fixture independently of openpyxl."""
    main = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    rel = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    package = "http://schemas.openxmlformats.org/package/2006/relationships"
    parts = {
        "[Content_Types].xml": (
            '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
            '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
            '<Default Extension="xml" ContentType="application/xml"/>'
            '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
            '<Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
            '<Override PartName="/xl/worksheets/sheet2.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
            '</Types>'
        ),
        "_rels/.rels": f'<Relationships xmlns="{package}"><Relationship Id="rId1" Type="{rel}/officeDocument" Target="xl/workbook.xml"/></Relationships>',
        "xl/workbook.xml": (
            f'<workbook xmlns="{main}" xmlns:r="{rel}"><sheets>'
            '<sheet name="Example" sheetId="1" r:id="rId1"/>'
            '<sheet name="Hidden" sheetId="2" state="hidden" r:id="rId2"/>'
            '</sheets></workbook>'
        ),
        "xl/_rels/workbook.xml.rels": (
            f'<Relationships xmlns="{package}">'
            f'<Relationship Id="rId1" Type="{rel}/worksheet" Target="worksheets/sheet1.xml"/>'
            f'<Relationship Id="rId2" Type="{rel}/worksheet" Target="worksheets/sheet2.xml"/>'
            '</Relationships>'
        ),
        "xl/worksheets/sheet1.xml": (
            f'<worksheet xmlns="{main}"><dimension ref="A1:B2"/><sheetData>'
            '<row r="1"><c r="A1" t="inlineStr"><is><t>SYNTHETIC</t></is></c><c r="B1"><v>7</v></c></row>'
            '<row r="2"><c r="A2"><v>3</v></c><c r="B2"><f>A2*2</f><v>6</v></c></row>'
            '</sheetData></worksheet>'
        ),
        "xl/worksheets/sheet2.xml": (
            f'<worksheet xmlns="{main}"><dimension ref="A1"/><sheetData>'
            '<row r="1"><c r="A1" t="inlineStr"><is><t>SYNTHETIC-HIDDEN</t></is></c></row>'
            '</sheetData></worksheet>'
        ),
    }
    with zipfile.ZipFile(path, "x", zipfile.ZIP_DEFLATED) as archive:
        for name, content in sorted(parts.items()):
            info = zipfile.ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            archive.writestr(info, content)
    os.chmod(path, 0o600)


def extract_cells(path: Path, *, data_only: bool) -> list[dict]:
    from openpyxl import load_workbook

    units = []
    with source_read_guard(path) as identity:
        workbook = load_workbook(path, read_only=True, data_only=data_only, keep_links=False)
        try:
            for sheet_index, sheet in enumerate(workbook.worksheets):
                for row in sheet.iter_rows():
                    for cell in row:
                        if cell.value is None:
                            continue
                        unit = {
                            "schema_version": 1, "run_id": "openpyxl-demo",
                            "sample_id": "SYN-XLSX-MODES", "source_kind": "synthetic",
                            "source_sha256": identity.sha256, "lane": "candidate",
                            "unit_id": f"cell:{sheet_index}:{cell.coordinate}", "unit_type": "cell",
                            "locator": {"sheet_index": sheet_index, "cell": cell.coordinate},
                            "content_hashes": {}, "status": "ok", "warnings": [],
                            "attributes": {"formula": cell.data_type == "f"},
                        }
                        validate_unit(unit)
                        units.append(unit)
        finally:
            workbook.close()
    return units


def run_demo(output_root: Path) -> dict:
    from openpyxl import __version__

    if output_root.exists() or output_root.is_symlink():
        raise FileExistsError("choose a new output directory")
    root = secure_directory(output_root)
    with deny_python_network() as attempts:
        source = root / "synthetic.xlsx"
        build_workbook(source)
        census = census_file(source)
        structure = census["structure"]
        expected = {
            "cells": sum(sheet["cell_nodes"] for sheet in structure["sheets"]),
            "formula_cells": sum(sheet["formula_cells"] for sheet in structure["sheets"]),
        }
        modes = {}
        for name, data_only in (("formulas", False), ("cached_values", True)):
            units = extract_cells(source, data_only=data_only)
            modes[name] = {
                "data_only": data_only, "read_only": True,
                "cells": len(units),
                "formula_cells": sum(unit["attributes"]["formula"] for unit in units),
                "locators": [unit["locator"] for unit in units],
            }
        report = {
            "schema_version": 1, "mode": "synthetic_openpyxl_reading_modes",
            "extractor": "openpyxl", "extractor_version": __version__,
            "source_sha256": census["source_sha256"], "source_identity_unchanged": True,
            "expected": expected, "modes": modes,
            "expected_tradeoff_observed": (
                expected["formula_cells"] > 0
                and modes["formulas"]["formula_cells"] == expected["formula_cells"]
                and modes["cached_values"]["formula_cells"] == 0
                and modes["formulas"]["cells"] == modes["cached_values"]["cells"] == expected["cells"]
                and modes["formulas"]["locators"] == modes["cached_values"]["locators"]
            ),
            "network_attempts": len(attempts), "private_sources_opened": 0,
            "limits": ["one_generated_workbook", "counts_and_cell_locators_only", "cached_values_not_recalculated"],
        }
    report["report_sha256"] = hashlib.sha256(
        json.dumps(report, sort_keys=True, separators=(",", ":")).encode()
    ).hexdigest()
    atomic_json(root / "comparison.json", report)
    return report


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, required=True)
    report = run_demo(parser.parse_args().output)
    print(json.dumps(report, indent=2, sort_keys=True))
    if not report["expected_tradeoff_observed"]:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
